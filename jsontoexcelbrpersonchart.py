import pandas as pd
import numpy as np
from openpyxl import Workbook
import csv
import os
import xlsxwriter
from xlsxwriter.utility import xl_rowcol_to_cell
from shutil import copyfile
from pathlib import Path
import json
from pandas.io.json import json_normalize

p = Path(r'C:\Data\json\backregisterperson.json')
with p.open('r', encoding='utf-8') as f:
    data = json.loads(f.read())
df = pd.json_normalize(data['Persons'])

#df = df.sort_values(["Person.Index"], axis=0, ascending=[True])
#https://thispointer.com/pandas-sort-rows-or-columns-in-dataframe-based-on-values-using-dataframe-sort_values/

df = df.sort_values(by =['Person.Index', 'Timestamp'])

#először a Personal index majd a Timestamp alapján rendezi sorba !

#print (df)

writer = pd.ExcelWriter('backregisterpersonchart.xlsx', engine='xlsxwriter')

df.to_excel(writer, sheet_name='Sheet1')

workbook = writer.book

worksheet = writer.sheets['Sheet1']

format1 = workbook.add_format({'num_format': '#,##0'}) #tizedesjegy beállítás

worksheet.set_column('V:V', 18, format1)

worksheet.set_column('X:X', 18, format1)

worksheet.write_formula('H2', '=COUNTIF(B2:B2180,"<10000")') #Fontos az excelben szereplő pontosvessző helyett ; sima kell ,

worksheet.write_formula('H3', '=COUNTIF(B2:B2180,"<20000")-H2')

worksheet.write_formula('H4', '=COUNTIF(B2:B2180,"<30000")-H2-H3')

worksheet.write_formula('H5', '=COUNTIF(B2:B2180,"<40000")-H2-H3-H4')

worksheet.write_formula('H6', '=COUNTIF(B2:B2180,"<50000")-H2-H3-H4-H5')

worksheet.write_formula('H7', '=COUNTIF(B2:B2180,"<60000")-H2-H3-H4-H5-H6')

worksheet.write('I2', '0:10')

worksheet.write('I3', '10:20')

worksheet.write('I4', '20:30')

worksheet.write('I5', '30:40')

worksheet.write('I6', '40:50')

worksheet.write('I7', '50:60')

worksheet.write_formula('H10', '=MAX(C2:C2178)')

worksheet.write_array_formula('K2:K2180', '{=IF($C$2:$C$2178=0,$B$2:$B$2178,"")}')

worksheet.write_array_formula('L2:L2180', '{=IF($C$2:$C$2178=1,$B$2:$B$2178,"")}')

worksheet.write_array_formula('M2:M2180', '{=IF($C$2:$C$2178=2,$B$2:$B$2178,"")}')

worksheet.write_array_formula('N2:N2180', '{=IF($C$2:$C$2178=3,$B$2:$B$2178,"")}')

worksheet.write_array_formula('O2:O2180', '{=IF($C$2:$C$2178=4,$B$2:$B$2178,"")}')

worksheet.write_array_formula('P2:P2180', '{=IF($C$2:$C$2178=5,$B$2:$B$2178,"")}')

worksheet.write_array_formula('Q2:Q2180', '{=IF($C$2:$C$2178=6,$B$2:$B$2178,"")}')

worksheet.write_array_formula('R2:R2180', '{=IF($C$2:$C$2178=7,$B$2:$B$2178,"")}')

worksheet.write_array_formula('S2:S2180', '{=IF($C$2:$C$2178=8,$B$2:$B$2178,"")}')

worksheet.write_array_formula('T2:T2180', '{=IF($C$2:$C$2178=9,$B$2:$B$2178,"")}')

worksheet.write_formula('V2', '=(MAX(K2:K2178)-MIN(K2:K2178))/1000')

worksheet.write_formula('V3', '=(MAX(L2:L2178)-MIN(L2:L2178))/1000')

worksheet.write_formula('V4', '=(MAX(M2:M2178)-MIN(M2:M2178))/1000')

worksheet.write_formula('V5', '=(MAX(N2:N2178)-MIN(N2:N2178))/1000')

worksheet.write_formula('V6', '=(MAX(O2:O2178)-MIN(O2:O2178))/1000')

worksheet.write_formula('V7', '=(MAX(P2:P2178)-MIN(P2:P2178))/1000')

worksheet.write_formula('V8', '=(MAX(Q2:Q2178)-MIN(Q2:Q2178))/1000')

worksheet.write_formula('V9', '=(MAX(R2:R2178)-MIN(R2:R2178))/1000')

worksheet.write_formula('V10', '=(MAX(S2:S2178)-MIN(S2:S2178))/1000')

worksheet.write_formula('V11', '=(MAX(T2:T2178)-MIN(T2:T2178))/1000')

worksheet.write_formula('X2', '=AVERAGE(V2:V11)')

worksheet.write_formula('X3', '=AVERAGE(V2:V11)')

worksheet.write_formula('X4', '=AVERAGE(V2:V11)')

worksheet.write_formula('X5', '=AVERAGE(V2:V11)')

worksheet.write_formula('X6', '=AVERAGE(V2:V11)')

worksheet.write_formula('X7', '=AVERAGE(V2:V11)')

worksheet.write_formula('X8', '=AVERAGE(V2:V11)')

worksheet.write_formula('X9', '=AVERAGE(V2:V11)')

worksheet.write_formula('X10', '=AVERAGE(V2:V11)')

worksheet.write_formula('X11', '=AVERAGE(V2:V11)')

worksheet.write('W2', '1')

worksheet.write('W3', '2')

worksheet.write('W4', '3')

worksheet.write('W5', '4')

worksheet.write('W6', '5')

worksheet.write('W7', '6')

worksheet.write('W8', '7')

worksheet.write('W9', '8')

worksheet.write('W10', '9')

worksheet.write('W11', '10')

worksheet.write('Y2', 'Átlag')

#Fontos az excelben szereplő pontosvessző helyett ; sima kell ,

worksheet = workbook.add_worksheet()

chart = workbook.add_chart({'type': 'column'})

chart.add_series({
    'name': ' ',
    'values': '=Sheet1!$V$2:$V$11',
    'categories': '=Sheet1!$W$2:$W$11',
    'data_labels': {'value': True, 'font':{'size': 16, 'bold': True}},
})

chart.set_title({
    'name': 'Az egyes vásárlók adott áruházi sorban eltöltött ideje és ennek átlaga',
    'overlay': True,
    'position': 'center',
    'layout': {
        'x': 0.17,
        'y': 0.03,
    },
    'name_font': {
        'name': 'Calibri',
        'bold': True,
        'color': 'black',
        'size': 28
    },
})

chart.set_x_axis({
    'name': 'Fő',
    'name_layout': {
        'x': 0.35,
        'y': 0.95,
    },
    'name_font': {
        'name': 'Calibri',
        'bold': True,
        'color': 'black',
        'size': 24
    },
    'num_font': {
        'name': 'Calibri',
        'bold': True,
        'color': 'black',
        'size': 16
    },
})

chart.set_y_axis({
    'name': 'Másodperc',
    'name_font': {
        'name': 'Calibri',
        'bold': True,
        'color': 'black',
        'size': 24
    },
    'num_font': {
        'name': 'Calibri',
        'bold': True,
        'color': 'black',
        'size': 16
    },
})

chart.set_legend({
    'none': False,
    'position': 'bottom',
    'layout': {
        'x':      0.80,
        'y':      0.40,
        'width':  0.12,
        'height': 0.25,
    }
})

chart.set_plotarea({
    'layout': {
        'x':      0.17,
        'y':      0.20,
        'width':  0.70,
        'height': 0.64,
    }
})

chart.set_style(11)

line_chart = workbook.add_chart({'type': 'line'})

line_chart.add_series({
    'name': ' ',
    'values': '=Sheet1!$X$2:$X$11',
    'categories': '=Sheet1!$Y$2',
    'line': {'color': 'red'},
})

line_chart.set_legend({
    'none': False,
    'position': 'bottom',
    'layout': {
        'x':      0.80,
        'y':      0.40,
        'width':  0.12,
        'height': 0.25,
    }
})

chart.combine(line_chart)

worksheet.insert_chart('C4', chart,  {'x_scale': 2, 'y_scale': 2})

worksheet = workbook.add_worksheet()

chart = workbook.add_chart({'type': 'column'})

chart.add_series({
    'name': ' ',
    'values': '=Sheet1!$H$2:$H$7',
    'categories': '=Sheet1!$I$2:$I$7',
    'data_labels': {'value': True, 'font':{'size': 16, 'bold': True}},
})

chart.set_title({
    'name': '10 másodpercenként detektált esetek 1 perc alatt',
    'overlay': True,
    'position': 'center',
    'layout': {
        'x': 0.17,
        'y': 0.03,
    },
    'name_font': {
        'name': 'Calibri',
        'bold': True,
        'color': 'black',
        'size': 28
    },
})

chart.set_x_axis({
    'name': 'Másodperc',
    'name_layout': {
        'x': 0.35,
        'y': 0.95,
    },
    'name_font': {
        'name': 'Calibri',
        'bold': True,
        'color': 'black',
        'size': 24
    },
    'num_font': {
        'name': 'Calibri',
        'bold': True,
        'color': 'black',
        'size': 16
    },
})

chart.set_y_axis({
    'name': 'db',
    'name_font': {
        'name': 'Calibri',
        'bold': True,
        'color': 'black',
        'size': 24
    },
    'num_font': {
        'name': 'Calibri',
        'bold': True,
        'color': 'black',
        'size': 16
    },
})

#chart.set_legend({'font': {'bold': 1, 'italic': 1}})

chart.set_legend({
    'none': True,
    'position': 'bottom',
    'layout': {
        'x':      0.80,
        'y':      0.40,
        'width':  0.12,
        'height': 0.25,
    }
})

chart.set_plotarea({
    'layout': {
        'x':      0.17,
        'y':      0.20,
        'width':  0.70,
        'height': 0.64,
    }
})

chart.set_style(11)

#https://xlsxwriter.readthedocs.io/working_with_charts.html

#https://xlsxwriter.readthedocs.io/example_chart_column.html

#https://pandas-xlsxwriter-charts.readthedocs.io/chart_examples.html

worksheet.insert_chart('C4', chart,  {'x_scale': 2, 'y_scale': 2}) #ezzel növelhető a terület

writer.save()

#https://xlsxwriter.readthedocs.io/example_pandas_chart.html
# F5-el kell elindítani akkor csinálja meg és a 3.7-el !!!
#https://xlsxwriter.readthedocs.io/worksheet.html